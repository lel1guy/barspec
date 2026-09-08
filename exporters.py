"""Export builders — the owner-facing data escape hatches.

Pure-ish functions taking db structures; routes in main.py are thin wrappers.
Costs ARE included in these exports on purpose: the buyer of these files is
the owner (or their accountant), not the floor staff.
"""
import io
import qrcode
import qrcode.image.svg

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HEAD_FILL = PatternFill("solid", fgColor="1F2430")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
    ws.freeze_panes = "A2"


def _autosize(ws) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[letter].width = min(width, 48)


def specs_workbook(specs_full: list[dict]) -> io.BytesIO:
    """Two sheets: the spec book + every ingredient line, per spec."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Specs"
    ws.append(["Name", "Category", "Glass", "Method", "Garnish",
               "Dilution %", "Served ml", "ABV %", "Cost €", "Sell €",
               "Margin %", "Target %", "Dietary", "Allergens"])
    for s in specs_full:
        sm = s["summary"]
        ws.append([
            s["name"], s.get("category") or "", s.get("glass", ""),
            s.get("method", ""), s.get("garnish", ""),
            sm.get("dilution_pct", 0), sm.get("served_ml", sm.get("total_ml", 0)),
            sm.get("served_abv", sm.get("abv", 0)), sm.get("cost_eur"),
            s.get("price_eur") if s.get("price_eur") else "",
            sm.get("margin"), s.get("target_gp"),
            s.get("dietary") or "", s.get("allergens") or "",
        ])
    _style_header(ws, ws.max_column)

    ws2 = wb.create_sheet("Ingredients")
    ws2.append(["Spec", "Ingredient", "Amount", "Unit", "Line cost €"])
    for s in specs_full:
        for l in s["lines"]:
            ws2.append([s["name"], l["name"], l["amount_ml"], l.get("unit") or "ml",
                        round(l["row_cost_eur"], 4)])
    _style_header(ws2, ws2.max_column)

    for sheet in (ws, ws2):
        _autosize(sheet)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def stock_workbook(items: list[dict]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Item", "Kind", "ABV %", "Price €", "Size", "Size unit",
               "Yield %", "Par", "Used in specs"])
    for it in items:
        ws.append([
            it["name"], it.get("dimension", "volume"), it.get("abv", 0),
            it.get("bottle_price_eur"), it.get("bottle_volume_ml"),
            {"volume": "ml", "weight": "g", "count": "pc"}.get(it.get("dimension"), "ml"),
            round((it.get("yield_frac") or 1) * 100), it.get("par_level") or "",
            it.get("used_in", 0),
        ])
    _style_header(ws, ws.max_column)
    _autosize(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def menu_qr_svg(target_url: str) -> str:
    """SVG QR of the menu link — no raster deps, prints/scales forever."""
    qr = qrcode.QRCode(border=2, box_size=8)
    qr.add_data(target_url)
    qr.make(fit=True)
    return qr.make_image(image_factory=qrcode.image.svg.SvgPathImage).to_string().decode()
