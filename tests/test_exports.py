"""Export endpoints (010): owner-facing .xlsx/.csv + the menu QR.

The escape hatches must stay honest: the spec book and stock sheet carry
costs on purpose (owner/accountant files), the QR opens the menu view.
"""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import main

client = TestClient(main.app)


@pytest.fixture(autouse=True)
def _fresh(fresh_db):
    yield


def test_specs_xlsx_two_sheets():
    r = client.get("/api/export/specs.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "barspec-specs.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Specs"]
    rows = list(ws.values)
    assert len(rows) == 6  # header + 5 seed specs
    assert rows[0][0] == "Name"
    assert rows[1][1] in (None, "")  # seed specs are uncategorised
    ing = wb["Ingredients"]
    assert len(list(ing.values)) == 16  # header + 15 spec lines
    assert "Negroni" in {r[0] for r in list(ing.values)[1:]}


def test_stock_xlsx():
    r = client.get("/api/export/stock.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb["Stock"].values)
    assert len(rows) == 16  # header + 15 seed items
    assert rows[0][0] == "Item"
    gin = [r for r in rows[1:] if r[0] == "London dry gin"][0]
    assert gin[2] == 40.0  # ABV


def test_specs_csv_rows_match():
    r = client.get("/api/export/specs.csv")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    lines = r.text.strip().splitlines()
    assert len(lines) == 6  # header + 5 specs
    assert lines[0].startswith("Name,Category")


def test_stock_csv_rows_match():
    r = client.get("/api/export/stock.csv")
    assert len(r.text.strip().splitlines()) == 16


def test_menu_qr_svg():
    r = client.get("/api/export/menu-qr.svg",
                   params={"url": "http://192.168.1.77:8777/?view=menu"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "image/svg+xml"
    assert "<svg" in r.text
    # SvgPathImage emits one path whose data carries every QR module
    d = r.text.split('d="')[1].split('"')[0]
    assert len(d) > 500  # real QR modules, not a blank image
